import json
import base64
from io import BytesIO


def parse_xlsx(data: bytes):
    import openpyxl
    wb = openpyxl.load_workbook(BytesIO(data), data_only=True)
    ws = wb.active
    rows = []
    for row in ws.iter_rows(values_only=True):
        cells = ['' if c is None else str(c).strip() for c in row]
        if any(cells):
            rows.append(cells)
    return rows


def parse_docx(data: bytes):
    import docx
    doc = docx.Document(BytesIO(data))
    rows = []
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                rows.append(cells)
    return rows


def handler(event: dict, context) -> dict:
    '''Разбирает входящий файл с товарами (.xlsx или .docx) и возвращает строки таблицы.'''
    if event.get('httpMethod') == 'OPTIONS':
        return {
            'statusCode': 200,
            'headers': {
                'Access-Control-Allow-Origin': '*',
                'Access-Control-Allow-Methods': 'POST, OPTIONS',
                'Access-Control-Allow-Headers': 'Content-Type',
                'Access-Control-Max-Age': '86400',
            },
            'body': '',
        }

    cors = {'Access-Control-Allow-Origin': '*', 'Content-Type': 'application/json'}
    body = json.loads(event.get('body') or '{}')
    filename = (body.get('filename') or '').lower()
    file_b64 = body.get('file', '')

    if not file_b64:
        return {'statusCode': 400, 'headers': cors, 'isBase64Encoded': False,
                'body': json.dumps({'error': 'Файл не передан'}, ensure_ascii=False)}

    data = base64.b64decode(file_b64)

    if filename.endswith('.xlsx'):
        rows = parse_xlsx(data)
        ftype = 'xlsx'
    elif filename.endswith('.docx'):
        rows = parse_docx(data)
        ftype = 'docx'
    else:
        return {'statusCode': 400, 'headers': cors, 'isBase64Encoded': False,
                'body': json.dumps({'error': 'Поддерживаются только .xlsx и .docx'}, ensure_ascii=False)}

    return {'statusCode': 200, 'headers': cors, 'isBase64Encoded': False,
            'body': json.dumps({'type': ftype, 'row_count': len(rows), 'rows': rows},
                               ensure_ascii=False)}
