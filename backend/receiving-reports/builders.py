import base64
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    Image, KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
)

from embedded import FILES


def _asset(name):
    return io.BytesIO(base64.b64decode(FILES[name]))

COLS = [('sale', 'На продажу'), ('wipe', 'На протирку'), ('repair', 'Под ремонт'), ('scrap', 'В утиль')]
TITLE = dict(COLS)
# Куда ушло при проверке → тот же склад и цвет, что в плашках ниже.
OUT_WH = {'sale': 'СГП', 'wipe': 'Протирка', 'repair': 'Под ремонт', 'scrap': 'Утиль'}

TILES = {
    'СГП': ('package-check.png', '34D399', 'E8FBF3'),
    'Протирка': ('sparkles.png', '38BDF8', 'E7F6FE'),
    'Под ремонт': ('wrench.png', 'F59E0B', 'FEF5E3'),
    'Утиль': ('trash-2.png', 'FB7185', 'FFECEF'),
}

ACCENT = '5B3FA8'


def build_xlsx(data, meta):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Сводка'
    thin = Side(style='thin', color='D0D0D0')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill('solid', fgColor=ACCENT)
    head_font = Font(bold=True, color='FFFFFF')
    sec_font = Font(bold=True, size=13, color=ACCENT)

    ws.column_dimensions['A'].width = 58
    ws.column_dimensions['B'].width = 18
    for col in 'CDEFG':
        ws.column_dimensions[col].width = 14

    r = 1
    ws.cell(r, 1, 'Сводка за период + содержание складов').font = Font(bold=True, size=15)
    r += 1
    ws.cell(r, 1, f"Период: {meta['period']}   ·   Мастер: {meta['master']}")
    r += 1
    ws.cell(r, 1, f"Сформирован: {meta['made']}").font = Font(color='777777', size=9)
    r += 2

    ws.cell(r, 1, '1. Итог за период').font = sec_font
    r += 1
    ws.cell(r, 1, 'Проверено единиц')
    c = ws.cell(r, 2, data['total'])
    c.font = Font(bold=True, size=14)
    c.alignment = Alignment(horizontal='center')
    r += 1
    for s in data['split']:
        _, color, bg = TILES[OUT_WH[s['key']]]
        a = ws.cell(r, 1, TITLE[s['key']])
        a.alignment = Alignment(indent=2)
        b = ws.cell(r, 2, s['qty'])
        b.alignment = Alignment(horizontal='center')
        if s['label'] == '<1%':
            c = ws.cell(r, 3, s['qty'] / data['total'])
            c.number_format = '"<1%"'
        else:
            c = ws.cell(r, 3, s['pct'] / 100)
            c.number_format = '0%'
        c.alignment = Alignment(horizontal='center')
        for x in (a, b, c):
            x.fill = PatternFill('solid', fgColor=bg)
            x.border = border
        a.border = Border(left=Side(style='thick', color=color), right=thin, top=thin, bottom=thin)
        b.font = Font(bold=True, color='333333')
        r += 1
    r += 1

    ws.cell(r, 1, '2. Проверенный товар за период').font = sec_font
    r += 1
    heads = ['Техническое наименование', 'Заказ-наряд'] + [t for _, t in COLS] + ['Всего']
    for i, h in enumerate(heads, 1):
        c = ws.cell(r, i, h)
        c.fill, c.font, c.border = head_fill, head_font, border
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    r += 1
    sums = {k: 0 for k, _ in COLS}
    for row in data['checked']:
        vals = [row['name'], row['order']] + [row[k] or None for k, _ in COLS] + [row['total']]
        for i, v in enumerate(vals, 1):
            c = ws.cell(r, i, v)
            c.border = border
            c.alignment = Alignment(wrap_text=(i == 1), vertical='top',
                                    horizontal='left' if i <= 2 else 'center')
        for k, _ in COLS:
            sums[k] += row[k]
        r += 1
    if not data['checked']:
        ws.cell(r, 1, 'За период ничего не проверено').font = Font(italic=True, color='777777')
        r += 1
    tot = ['Итого', ''] + [sums[k] for k, _ in COLS] + [data['total']]
    for i, v in enumerate(tot, 1):
        c = ws.cell(r, i, v)
        c.font, c.border = Font(bold=True), border
        c.fill = PatternFill('solid', fgColor='EEEAF7')
        c.alignment = Alignment(horizontal='left' if i <= 2 else 'center')
    r += 2

    ws.cell(r, 1, '3. Склады сейчас').font = sec_font
    r += 1
    for w in data['warehouses']:
        _, color, bg = TILES[w['name']]
        a = ws.cell(r, 1, w['name'])
        b = ws.cell(r, 2, w['qty'])
        for c in (a, b):
            c.fill = PatternFill('solid', fgColor=bg)
            c.border = Border(left=Side(style='thick', color=color) if c is a else thin,
                              right=thin, top=thin, bottom=thin)
            c.font = Font(bold=True, size=12, color='333333')
        b.alignment = Alignment(horizontal='center')
        r += 1
    r += 1

    ws.cell(r, 1, '4. Что лежит на складах сейчас').font = sec_font
    r += 1
    for w in data['warehouses']:
        _, color, bg = TILES[w['name']]
        a = ws.cell(r, 1, w['name'])
        b = ws.cell(r, 2, w['qty'])
        for c in (a, b):
            c.fill = PatternFill('solid', fgColor=bg)
            c.font = Font(bold=True, size=12)
        r += 1
        if not w['dirs']:
            c = ws.cell(r, 1, 'пусто')
            c.font = Font(italic=True, color='999999')
            c.alignment = Alignment(indent=2)
            r += 1
        for d in w['dirs']:
            c = ws.cell(r, 1, d['name'])
            c.font = Font(bold=True)
            c.alignment = Alignment(indent=2)
            ws.cell(r, 2, d['qty']).font = Font(bold=True)
            r += 1
            for name, n in d['items']:
                c = ws.cell(r, 1, name)
                c.alignment = Alignment(indent=5, wrap_text=True)
                ws.cell(r, 2, n)
                r += 1
        r += 1

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


_fonts_ready = False


def _fonts():
    global _fonts_ready
    if not _fonts_ready:
        pdfmetrics.registerFont(TTFont('DV', _asset('DejaVuSans.ttf')))
        pdfmetrics.registerFont(TTFont('DVB', _asset('DejaVuSans-Bold.ttf')))
        _fonts_ready = True


def _hex(h):
    return colors.HexColor('#' + h)


def build_pdf(data, meta):
    _fonts()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=landscape(A4),
        leftMargin=12 * mm, rightMargin=12 * mm, topMargin=12 * mm, bottomMargin=12 * mm,
        title='Сводка за период + содержание складов',
    )
    width = doc.width
    st = {
        'title': ParagraphStyle('t', fontName='DVB', fontSize=16, leading=20),
        'meta': ParagraphStyle('m', fontName='DV', fontSize=10, leading=14),
        'small': ParagraphStyle('s', fontName='DV', fontSize=8, leading=10, textColor=colors.grey),
        'sec': ParagraphStyle('h', fontName='DVB', fontSize=13, leading=16,
                              textColor=_hex(ACCENT), spaceBefore=8, spaceAfter=6),
        'cell': ParagraphStyle('c', fontName='DV', fontSize=8.5, leading=10.5),
        'cellb': ParagraphStyle('cb', fontName='DVB', fontSize=8.5, leading=10.5),
        'head': ParagraphStyle('hd', fontName='DVB', fontSize=8.5, leading=10.5,
                               textColor=colors.white, alignment=1),
        'wh': ParagraphStyle('w', fontName='DVB', fontSize=12, leading=15),
        'dir': ParagraphStyle('d', fontName='DVB', fontSize=10, leading=13, leftIndent=8 * mm),
        'item': ParagraphStyle('i', fontName='DV', fontSize=9, leading=12, leftIndent=18 * mm),
        'empty': ParagraphStyle('e', fontName='DV', fontSize=9, leading=12, leftIndent=8 * mm,
                                textColor=colors.grey),
    }

    def p(text, style):
        return Paragraph(str(text).replace('&', '&amp;').replace('<', '&lt;'), st[style])

    story = [
        p('Сводка за период + содержание складов', 'title'),
        Spacer(1, 2 * mm),
        p(f"Период: {meta['period']}   ·   Мастер: {meta['master']}", 'meta'),
        p(f"Сформирован: {meta['made']}", 'small'),
        Spacer(1, 4 * mm),
        p('1. Итог за период', 'sec'),
    ]
    total_box = Table(
        [[p('Проверено единиц', 'meta'),
          Paragraph(f"{data['total']}", ParagraphStyle('big', fontName='DVB', fontSize=22, leading=26))]],
        colWidths=[60 * mm, 40 * mm],
        hAlign='LEFT',
    )
    total_box.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), _hex('EEEAF7')),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('LEFTPADDING', (0, 0), (-1, -1), 8), ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
    ]))
    split_rows = []
    for s in data['split']:
        icon, color, bg = TILES[OUT_WH[s['key']]]
        split_rows.append([
            Image(_asset(icon), 5.5 * mm, 5.5 * mm),
            Paragraph(TITLE[s['key']], ParagraphStyle('sn', fontName='DV', fontSize=10, leading=13)),
            Paragraph(str(s['qty']), ParagraphStyle('sq', fontName='DVB', fontSize=11, leading=13,
                                                    alignment=2, textColor=_hex(color))),
            Paragraph(s['label'], ParagraphStyle('sp', fontName='DV', fontSize=10, leading=13,
                                                 alignment=2, textColor=_hex('555555'))),
        ])
    split = Table(split_rows, colWidths=[9 * mm, 45 * mm, 22 * mm, 18 * mm], hAlign='LEFT')
    style = [
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 3), ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
    ]
    for i, s in enumerate(data['split']):
        _, color, bg = TILES[OUT_WH[s['key']]]
        style += [('BACKGROUND', (0, i), (-1, i), _hex(bg)),
                  ('LINEBEFORE', (0, i), (0, i), 3, _hex(color)),
                  ('LINEBELOW', (0, i), (-1, i), 1.5, colors.white)]
    split.setStyle(TableStyle(style))
    indent = Table([['', split]], colWidths=[8 * mm, 96 * mm], hAlign='LEFT')
    indent.setStyle(TableStyle([('LEFTPADDING', (0, 0), (-1, -1), 0),
                                ('TOPPADDING', (0, 0), (-1, -1), 3)]))
    story += [total_box, indent, p('2. Проверенный товар за период', 'sec')]

    num_w = 22 * mm
    name_w = width - 35 * mm - num_w * 5
    rows = [[p(h, 'head') for h in
             ['Техническое наименование', 'Заказ-наряд'] + [t for _, t in COLS] + ['Всего']]]
    sums = {k: 0 for k, _ in COLS}
    for row in data['checked']:
        rows.append([p(row['name'], 'cell'), p(row['order'], 'cell')]
                    + [str(row[k] or '') for k, _ in COLS] + [str(row['total'])])
        for k, _ in COLS:
            sums[k] += row[k]
    if not data['checked']:
        rows.append([p('За период ничего не проверено', 'cell'), '', '', '', '', '', ''])
    rows.append([p('Итого', 'cellb'), ''] + [str(sums[k]) for k, _ in COLS] + [str(data['total'])])
    t = Table(rows, colWidths=[name_w, 35 * mm] + [num_w] * 5, repeatRows=1)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), _hex(ACCENT)),
        ('GRID', (0, 0), (-1, -1), 0.4, _hex('D0D0D0')),
        ('FONTNAME', (0, 1), (-1, -1), 'DV'),
        ('FONTSIZE', (0, 1), (-1, -1), 8.5),
        ('ALIGN', (2, 1), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -2), [colors.white, _hex('F7F7F9')]),
        ('BACKGROUND', (0, -1), (-1, -1), _hex('EEEAF7')),
        ('FONTNAME', (0, -1), (-1, -1), 'DVB'),
    ]))
    story += [t, p('3. Склады сейчас', 'sec')]

    tiles = []
    tile_w = (width - 9 * mm) / 4
    for w in data['warehouses']:
        icon, color, bg = TILES[w['name']]
        inner = Table(
            [[Image(_asset(icon), 12 * mm, 12 * mm),
              [Paragraph(w['name'], ParagraphStyle('tn', fontName='DV', fontSize=10, leading=12,
                                                   textColor=_hex('555555'))),
               Paragraph(f"{w['qty']} шт.", ParagraphStyle('tq', fontName='DVB', fontSize=18,
                                                            leading=22, textColor=_hex(color)))]]],
            colWidths=[16 * mm, tile_w - 18 * mm],
        )
        inner.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), _hex(bg)),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('LINEBEFORE', (0, 0), (0, 0), 3, _hex(color)),
            ('TOPPADDING', (0, 0), (-1, -1), 7), ('BOTTOMPADDING', (0, 0), (-1, -1), 7),
        ]))
        tiles.append(inner)
    row = Table([tiles], colWidths=[tile_w + 2.25 * mm] * 4)
    row.setStyle(TableStyle([('LEFTPADDING', (0, 0), (-1, -1), 0),
                             ('RIGHTPADDING', (0, 0), (-1, -1), 3 * mm)]))
    story += [row, p('4. Что лежит на складах сейчас', 'sec')]

    for w in data['warehouses']:
        _, color, bg = TILES[w['name']]
        head = Table([[p(w['name'], 'wh'), p(f"{w['qty']} шт.", 'wh')]],
                     colWidths=[width - 35 * mm, 35 * mm])
        head.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), _hex(bg)),
            ('LINEBEFORE', (0, 0), (0, 0), 3, _hex(color)),
            ('ALIGN', (1, 0), (1, 0), 'RIGHT'),
        ]))
        block = [head, Spacer(1, 1.5 * mm)]
        if not w['dirs']:
            block.append(p('пусто', 'empty'))
        story.append(KeepTogether(block))
        for d in w['dirs']:
            lines = [[p(d['name'], 'dir'), p(str(d['qty']), 'cellb')]]
            lines += [[p(name, 'item'), p(str(n), 'cell')] for name, n in d['items']]
            dt = Table(lines, colWidths=[width - 35 * mm, 35 * mm])
            dt.setStyle(TableStyle([
                ('ALIGN', (1, 0), (1, -1), 'RIGHT'),
                ('LINEBELOW', (0, 0), (-1, 0), 0.4, _hex('D0D0D0')),
                ('TOPPADDING', (0, 0), (-1, -1), 1.5), ('BOTTOMPADDING', (0, 0), (-1, -1), 1.5),
            ]))
            story.append(dt)
        story.append(Spacer(1, 4 * mm))

    doc.build(story)
    return buf.getvalue()
