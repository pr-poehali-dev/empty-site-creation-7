"""Прайс-лист для оптовика: сборка по заявкам или по периоду и фирмам"""
import json
import os
import base64
import io
from datetime import datetime
import psycopg2
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

TEMP_PRODUCT_ID = 19
ALLOWED_ROLES = ['Управляющий']
CORS = {
    'Access-Control-Allow-Origin': '*',
    'Access-Control-Allow-Methods': 'GET, POST, OPTIONS',
    'Access-Control-Allow-Headers': 'Content-Type, Authorization, X-Authorization',
    'Access-Control-Max-Age': '86400',
}


def get_db():
    return psycopg2.connect(os.environ['DATABASE_URL'])


def get_user_by_token(cur, token):
    cur.execute(
        """SELECT u.id, u.phone, u.role FROM users u
           JOIN user_sessions s ON s.user_id = u.id
           WHERE s.token = %s AND s.expires_at > NOW()""",
        (token,)
    )
    return cur.fetchone()


def get_manager_info(cur, phone):
    cur.execute(
        """SELECT m.id, m.role_id, r.name FROM managers m
           LEFT JOIN roles r ON r.id = m.role_id
           WHERE m.phone = %s AND m.status = 'authorized'""",
        (phone,)
    )
    return cur.fetchone()


def json_resp(status, body):
    return {
        'statusCode': status,
        'headers': {**CORS, 'Content-Type': 'application/json'},
        'body': json.dumps(body, ensure_ascii=False, default=str),
    }


def esc(value):
    return str(value).replace("'", "''")


def collect_order_ids(cur, params):
    ids = params.get('order_ids') or []
    include_archived = bool(params.get('include_archived'))
    if ids:
        safe = [int(i) for i in ids]
        placeholders = ','.join(str(i) for i in safe)
        status_cond = '' if include_archived else " AND o.status <> 'archived'"
        cur.execute(
            f"SELECT o.id FROM wholesale_orders o "
            f"WHERE o.id IN ({placeholders}) AND o.status <> 'draft'{status_cond}"
        )
        return [r[0] for r in cur.fetchall()]

    conditions = ["o.status <> 'draft'"]
    if not include_archived:
        conditions.append("o.status <> 'archived'")

    date_from = params.get('date_from')
    date_to = params.get('date_to')
    if date_from:
        conditions.append(f"o.created_at >= '{esc(date_from)}'::date")
    if date_to:
        conditions.append(f"o.created_at < '{esc(date_to)}'::date + INTERVAL '1 day'")

    firms = params.get('firms') or []
    if firms:
        names = ','.join(f"'{esc(f)}'" for f in firms)
        conditions.append(f"o.customer_name IN ({names})")

    cur.execute(
        f"SELECT o.id FROM wholesale_orders o WHERE {' AND '.join(conditions)}"
    )
    return [r[0] for r in cur.fetchall()]


def build_price(cur, order_ids):
    if not order_ids:
        return [], []

    placeholders = ','.join(str(int(i)) for i in order_ids)
    cur.execute(
        f"""SELECT oi.product_id, COALESCE(oi.item_name, p.name), p.article, oi.price,
                   oi.temp_product_id, p.brand,
                   (SELECT pb.barcode FROM product_barcodes pb
                     WHERE pb.product_id = oi.product_id LIMIT 1),
                   tp.brand, tp.article, tp.nomenclature_id,
                   np.name, np.article, np.brand
            FROM wholesale_order_items oi
            JOIN products p ON p.id = oi.product_id
            LEFT JOIN temp_products tp ON tp.id = oi.temp_product_id
            LEFT JOIN products np ON np.id = tp.nomenclature_id
            WHERE oi.order_id IN ({placeholders})"""
    )

    merged = {}
    for row in cur.fetchall():
        (product_id, item_name, art, price, temp_id, brand, barcode,
         tp_brand, tp_article, tp_nom_id, np_name, np_article, np_brand) = row

        is_temp = temp_id is not None or product_id == TEMP_PRODUCT_ID
        if is_temp:
            if tp_nom_id and np_name:
                name, article, disp_brand = np_name, np_article or '', np_brand or ''
            elif tp_brand or tp_article:
                name = f"{tp_brand or ''} {tp_article or ''}".strip()
                article, disp_brand = tp_article or '', tp_brand or ''
            else:
                name = item_name or ''
                article = '' if art == '__TEMP__' else (art or '')
                disp_brand = ''
        else:
            name = item_name or ''
            article = '' if art == '__TEMP__' else (art or '')
            disp_brand = brand or ''

        key = (article.strip().lower(), name.strip().lower()) if article else ('', name.strip().lower())
        value = float(price or 0)

        if key in merged:
            if value > merged[key]['price']:
                merged[key]['price'] = value
            if not merged[key]['barcode'] and barcode:
                merged[key]['barcode'] = barcode
            if not merged[key]['brand'] and disp_brand:
                merged[key]['brand'] = disp_brand
        else:
            merged[key] = {
                'brand': disp_brand,
                'barcode': barcode or '',
                'article': article,
                'name': name,
                'price': value,
                'is_temp': is_temp,
            }

    all_items = list(merged.values())
    priced = sorted(
        [i for i in all_items if i['price'] > 0],
        key=lambda x: (x['brand'].lower(), x['name'].lower())
    )
    zero = sorted(
        [i for i in all_items if i['price'] <= 0],
        key=lambda x: (x['brand'].lower(), x['name'].lower())
    )
    return priced, zero


def make_excel(priced, zero):
    wb = Workbook()
    ws = wb.active
    ws.title = "Прайс-лист"

    thin = Side(style='thin')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_font = Font(bold=True, size=11)
    header_fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')

    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['B'].width = 18
    ws.column_dimensions['C'].width = 20
    ws.column_dimensions['D'].width = 50
    ws.column_dimensions['E'].width = 14

    ws['A1'] = "Прайс-лист"
    ws['A1'].font = Font(bold=True, size=13)
    ws['A2'] = f"Сформирован {datetime.now().strftime('%d.%m.%Y')}"
    ws['A2'].font = Font(size=10, color='666666')

    titles = ['Бренд', 'Штрихкод', 'Артикул', 'Наименование', 'Цена']
    start_row = 4
    for col_idx, title in enumerate(titles, 1):
        cell = ws.cell(row=start_row, column=col_idx, value=title)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal='center')

    row = start_row
    for item in priced:
        row += 1
        ws.cell(row=row, column=1, value=item['brand']).border = border
        ws.cell(row=row, column=2, value=item['barcode']).border = border
        ws.cell(row=row, column=3, value=item['article']).border = border
        ws.cell(row=row, column=4, value=item['name']).border = border
        price_cell = ws.cell(row=row, column=5, value=item['price'])
        price_cell.border = border
        price_cell.number_format = '#,##0.00'

    if zero:
        row += 2
        note = ws.cell(row=row, column=1, value='БЕЗ ЦЕНЫ — требуется проставить вручную')
        note.font = Font(bold=True, size=11, color='C00000')
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        row += 1
        for col_idx, title in enumerate(titles, 1):
            cell = ws.cell(row=row, column=col_idx, value=title)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = border
            cell.alignment = Alignment(horizontal='center')
        for item in zero:
            row += 1
            ws.cell(row=row, column=1, value=item['brand']).border = border
            ws.cell(row=row, column=2, value=item['barcode']).border = border
            ws.cell(row=row, column=3, value=item['article']).border = border
            ws.cell(row=row, column=4, value=item['name']).border = border
            ws.cell(row=row, column=5, value='').border = border

    buf = io.BytesIO()
    wb.save(buf)
    return base64.b64encode(buf.getvalue()).decode()


def handler(event: dict, context) -> dict:
    """Собирает прайс-лист для оптовика по выбранным заявкам или по периоду и фирмам"""
    method = event.get('httpMethod', 'GET')
    if method == 'OPTIONS':
        return {'statusCode': 200, 'headers': CORS, 'body': ''}

    req_headers = event.get('headers') or {}
    auth = req_headers.get('X-Authorization', '') or req_headers.get('Authorization', '')
    token = auth.replace('Bearer ', '').strip()

    conn = get_db()
    cur = conn.cursor()

    user = get_user_by_token(cur, token)
    if not user:
        cur.close(); conn.close()
        return json_resp(401, {'error': 'Не авторизован'})

    _user_id, phone, role = user
    is_owner = role == 'owner'
    if not is_owner:
        mgr = get_manager_info(cur, phone)
        if not mgr or mgr[2] not in ALLOWED_ROLES:
            cur.close(); conn.close()
            return json_resp(403, {'error': 'Нет доступа'})

    if method == 'GET':
        cur.execute(
            """SELECT DISTINCT o.customer_name FROM wholesale_orders o
               WHERE o.customer_name IS NOT NULL AND o.customer_name <> ''
                 AND o.status <> 'draft'
               ORDER BY o.customer_name"""
        )
        firms = [r[0] for r in cur.fetchall()]
        cur.close(); conn.close()
        return json_resp(200, {'firms': firms})

    body = json.loads(event.get('body') or '{}')
    order_ids = collect_order_ids(cur, body)
    priced, zero = build_price(cur, order_ids)

    if body.get('format') == 'xlsx':
        file_b64 = make_excel(priced, zero)
        cur.close(); conn.close()
        return json_resp(200, {
            'file': file_b64,
            'filename': f"Прайс-лист_{datetime.now().strftime('%d-%m-%Y')}.xlsx",
        })

    cur.close(); conn.close()
    return json_resp(200, {
        'items': priced,
        'zero_items': zero,
        'orders_count': len(order_ids),
        'total': len(priced) + len(zero),
    })
